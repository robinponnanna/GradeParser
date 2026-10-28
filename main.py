import os
import io
import re
import sys
import shutil
import platform
import subprocess
from collections import defaultdict
import pymupdf
from PIL import Image
import pytesseract
import pandas as pd
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


def select_pdf_files() -> list[str]:
    """
    Open the native file application manager depending on the operating system
    (macOS, Windows, Linux with GNOME) to let the user select one or more PDF files.
    """
    current_os = platform.system()
    selected_files = []

    print(f"Opening native file manager for {current_os}...")

    if current_os == "Linux":
        # GNOME uses zenity for native GTK file selection dialogs
        if shutil.which("zenity"):
            try:
                result = subprocess.run(
                    [
                        "zenity",
                        "--file-selection",
                        "--multiple",
                        "--separator=|",
                        "--file-filter=PDF Files | *.pdf *.PDF",
                        "--title=Select PDF Grade Card(s)"
                    ],
                    capture_output=True,
                    text=True
                )
                if result.returncode == 0 and result.stdout.strip():
                    selected_files = [f.strip() for f in result.stdout.strip().split("|") if f.strip()]
            except Exception as err:
                print(f"Zenity dialog error: {err}")

    elif current_os == "Darwin":  # macOS
        # Native macOS NSOpenPanel via AppleScript
        script = """
        set selectedFiles to choose file with prompt "Select PDF Grade Card(s)" of type {"pdf"} with multiple selections allowed
        set posixPaths to {}
        repeat with aFile in selectedFiles
            copy POSIX path of aFile to end of posixPaths
        end repeat
        set AppleScript's text item delimiters to linefeed
        return posixPaths as text
        """
        try:
            result = subprocess.run(["osascript", "-e", script], capture_output=True, text=True)
            if result.returncode == 0 and result.stdout.strip():
                selected_files = [f.strip() for f in result.stdout.strip().splitlines() if f.strip()]
        except Exception as err:
            print(f"macOS dialog error: {err}")

    elif current_os == "Windows":
        # Native Windows Common File Dialog via PowerShell
        ps_command = (
            "[System.Reflection.Assembly]::LoadWithPartialName('System.windows.forms') | Out-Null; "
            "$dialog = New-Object System.Windows.Forms.OpenFileDialog; "
            "$dialog.Filter = 'PDF Files (*.pdf)|*.pdf|All Files (*.*)|*.*'; "
            "$dialog.Multiselect = $true; "
            "$dialog.Title = 'Select PDF Grade Card(s)'; "
            "if ($dialog.ShowDialog() -eq [System.Windows.Forms.DialogResult]::OK) { "
            "$dialog.FileNames -join '|' "
            "}"
        )
        try:
            result = subprocess.run(
                ["powershell", "-NoProfile", "-Command", ps_command],
                capture_output=True,
                text=True
            )
            if result.returncode == 0 and result.stdout.strip():
                selected_files = [f.strip() for f in result.stdout.strip().split("|") if f.strip()]
        except Exception as err:
            print(f"Windows dialog error: {err}")

    # Fallback: Tkinter file dialog if native tool did not return files
    if not selected_files:
        try:
            import tkinter as tk
            from tkinter import filedialog
            root = tk.Tk()
            root.withdraw()
            root.attributes("-topmost", True)
            files = filedialog.askopenfilenames(
                title="Select PDF Grade Card(s)",
                filetypes=[("PDF Files", "*.pdf"), ("All Files", "*.*")]
            )
            root.destroy()
            if files:
                selected_files = list(files)
        except Exception:
            pass

    return selected_files


def get_save_file_path(default_name: str = "student_grades.xlsx") -> str:
    """
    Open the native OS save file manager dialog (macOS, Windows, Linux with GNOME)
    allowing the user to choose the destination folder and name the output Excel file.
    """
    current_os = platform.system()
    save_path = ""

    print(f"Opening native save file manager for {current_os}...")

    if current_os == "Linux":
        # GNOME uses zenity for GTK save file dialogs
        if shutil.which("zenity"):
            try:
                result = subprocess.run(
                    [
                        "zenity",
                        "--file-selection",
                        "--save",
                        "--confirm-overwrite",
                        f"--filename={default_name}",
                        "--file-filter=Excel Spreadsheet (*.xlsx) | *.xlsx",
                        "--title=Save Grade Details Excel File"
                    ],
                    capture_output=True,
                    text=True
                )
                if result.returncode == 0 and result.stdout.strip():
                    save_path = result.stdout.strip()
            except Exception as err:
                print(f"Zenity save dialog error: {err}")

    elif current_os == "Darwin":  # macOS
        # Native macOS save dialog via AppleScript
        script = f"""
        set chosenFile to choose file name with prompt "Save Grade Details Excel File" default name "{default_name}"
        return POSIX path of chosenFile
        """
        try:
            result = subprocess.run(["osascript", "-e", script], capture_output=True, text=True)
            if result.returncode == 0 and result.stdout.strip():
                save_path = result.stdout.strip()
        except Exception as err:
            print(f"macOS save dialog error: {err}")

    elif current_os == "Windows":
        # Native Windows SaveFileDialog via PowerShell
        ps_command = f"""
        [System.Reflection.Assembly]::LoadWithPartialName('System.windows.forms') | Out-Null;
        $dialog = New-Object System.Windows.Forms.SaveFileDialog;
        $dialog.Filter = 'Excel Files (*.xlsx)|*.xlsx|All Files (*.*)|*.*';
        $dialog.FileName = '{default_name}';
        $dialog.Title = 'Save Grade Details Excel File';
        if ($dialog.ShowDialog() -eq [System.Windows.Forms.DialogResult]::OK) {{
            $dialog.FileName
        }}
        """
        try:
            result = subprocess.run(
                ["powershell", "-NoProfile", "-Command", ps_command],
                capture_output=True,
                text=True
            )
            if result.returncode == 0 and result.stdout.strip():
                save_path = result.stdout.strip()
        except Exception as err:
            print(f"Windows save dialog error: {err}")

    # Fallback: Tkinter save file dialog
    if not save_path:
        try:
            import tkinter as tk
            from tkinter import filedialog
            root = tk.Tk()
            root.withdraw()
            root.attributes("-topmost", True)
            chosen = filedialog.asksaveasfilename(
                title="Save Grade Details Excel File",
                initialfile=default_name,
                defaultextension=".xlsx",
                filetypes=[("Excel Files", "*.xlsx"), ("All Files", "*.*")]
            )
            root.destroy()
            if chosen:
                save_path = chosen
        except Exception:
            pass

    if save_path and not save_path.lower().endswith(".xlsx"):
        save_path += ".xlsx"

    return save_path


def extract_student_info(full_text: str) -> tuple[str, str]:
    """
    Dynamically extract Student Name and Roll Number from document text
    based on document labels.
    """
    name = "Unknown"
    # Match "Student Name: <name>" or "Name: <name>"
    m_name = re.search(r"(?:Student\s+Name|Name)\s*:\s*([^\n\r]+)", full_text, re.IGNORECASE)
    if m_name:
        name = m_name.group(1).strip()
    else:
        # Match label on one line, value on the subsequent line
        lines = [l.strip() for l in full_text.splitlines() if l.strip()]
        for idx, line in enumerate(lines):
            if line.lower() in ("name", "student name") and idx + 1 < len(lines):
                candidate = lines[idx + 1]
                if not re.search(r"^(?:father|mother|uid|roll|sl|reg)", candidate, re.IGNORECASE):
                    name = candidate
                    break

    roll_no = "Unknown"
    # Match "Roll Number: <roll>" or "Roll No.: <roll>" or "Reg No.: <roll>"
    m_roll = re.search(r"(?:Roll\s*(?:Number|No\.?)|Reg(?:istration)?\s*(?:Number|No\.?))\s*[:\s]*([A-Z0-9]+)", full_text, re.IGNORECASE)
    if m_roll:
        roll_no = m_roll.group(1).strip()
    else:
        lines = [l.strip() for l in full_text.splitlines() if l.strip()]
        for idx, line in enumerate(lines):
            if line.lower() in ("roll no.", "roll no", "roll number", "reg no.", "reg no") and idx + 1 < len(lines):
                candidate = lines[idx + 1]
                if re.match(r"^[A-Z0-9]+$", candidate, re.IGNORECASE):
                    roll_no = candidate
                    break

    return name, roll_no


def extract_semester_info(full_text: str) -> str:
    """
    Dynamically fetch the Grade Card heading indicating when the examination was conducted
    directly from the document text, without hardcoding exam types or normalization dictionaries.
    """
    lines = [l.strip() for l in full_text.splitlines() if l.strip()]
    if not lines:
        return "N/A"

    # 1. On Grade Cards with "GRADE CARD", the line directly preceding it is the examination heading
    for idx, line in enumerate(lines):
        if line.upper() == "GRADE CARD" and idx > 0:
            return re.sub(r"\s+", " ", lines[idx - 1]).strip()

    # 2. On Provisional Grade Cards, line 1 contains the program/semester heading
    if lines[0].lower().startswith("provisional grade card") and len(lines) > 1:
        sem_heading = lines[1]
        sem_part = re.split(r"\s+b\.?tech", sem_heading, flags=re.IGNORECASE)[0].strip()
        # Look for the examination session period (Month Year) in the header section
        for cand in lines[2:16]:
            if re.search(r"\b(January|February|March|April|May|June|July|August|September|October|November|December)\s+\d{4}\b", cand, re.IGNORECASE):
                return f"{sem_part} ({cand})"
        return sem_part

    # 3. Dynamic search in header lines for examination / semester information
    for line in lines[:5]:
        if any(keyword in line.lower() for keyword in ["examination", "semester", "term"]):
            return re.sub(r"\s+", " ", line).strip()

    return lines[0]


def extract_result_date(full_text: str) -> str:
    """
    Dynamically extract the result publication / announcement date from document text.
    """
    date_match = re.search(
        r"(?:Result\s+Date|Date\s+of\s+Declaration|Date)\s*:\s*([^\n\r]+)",
        full_text,
        re.IGNORECASE
    )
    if date_match:
        date_str = date_match.group(1).strip()
        # Extract date pattern (e.g. DD/MM/YYYY, DD-MM-YYYY, or DD Month YYYY)
        valid_date = re.search(
            r"\b(?:\d{1,2}[\/\-]\d{1,2}[\/\-]\d{2,4}|\d{1,2}\s+[A-Za-z]+\s+\d{4})\b",
            date_str
        )
        if valid_date:
            return valid_date.group(0).strip()
        return date_str
    return "N/A"


def extract_page_words_with_ocr(page) -> tuple[str, list]:
    """
    Fallback OCR using pytesseract for scanned / image-based PDF pages.
    Renders the page to a high-res image and extracts text and word coordinates.
    Returns (page_text, words_list).
    words_list format: [(x0, y0, x1, y1, word), ...]
    """
    try:
        # Render high-resolution page image (300 DPI) for accurate recognition
        pix = page.get_pixmap(dpi=300)
        img = Image.open(io.BytesIO(pix.tobytes("png")))

        page_text = pytesseract.image_to_string(img)
        ocr_data = pytesseract.image_to_data(img, output_type=pytesseract.Output.DICT)

        words = []
        n_boxes = len(ocr_data.get("text", []))
        for i in range(n_boxes):
            word_str = ocr_data["text"][i].strip()
            if word_str:
                x0 = float(ocr_data["left"][i])
                y0 = float(ocr_data["top"][i])
                x1 = x0 + float(ocr_data["width"][i])
                y1 = y0 + float(ocr_data["height"][i])
                words.append((x0, y0, x1, y1, word_str))

        return page_text, words

    except pytesseract.TesseractNotFoundError:
        print("  [Notice] Tesseract OCR executable not found on system PATH. Install Tesseract to enable OCR for scanned PDFs.")
        return "", []
    except Exception as err:
        print(f"  [Notice] OCR extraction failed: {err}")
        return "", []


def parse_course_rows_from_words(words: list, line_tolerance: float = 6.0) -> list[tuple[str, str, int, str]]:
    """
    Dynamically group words by horizontal lines, identify course rows,
    and extract: (course_code, course_name, credits, grade) directly
    without relying on hardcoded grade whitelists or rigid course prefixes.
    """
    records = []
    # Alphanumeric course code: starts with letters/digits, contains letters and numbers
    course_code_regex = re.compile(r"^[A-Z0-9]{2,}[0-9]{2,}[A-Z0-9]*$", re.IGNORECASE)

    # Group words by vertical position (y-coordinate)
    lines_by_y = defaultdict(list)
    for w in words:
        y_mid = (w[1] + w[3]) / 2
        matched_key = None
        for key in lines_by_y:
            if abs(key - y_mid) < line_tolerance:
                matched_key = key
                break
        if matched_key is None:
            matched_key = y_mid
        lines_by_y[matched_key].append(w)

    # Sort lines top-to-bottom
    for y in sorted(lines_by_y.keys()):
        # Sort words left-to-right
        line_words = sorted(lines_by_y[y], key=lambda x: x[0])
        tokens = [x[4].strip() for x in line_words if x[4].strip()]

        # Find course code token in line
        code_idx = -1
        for idx, token in enumerate(tokens):
            if course_code_regex.match(token) and not token.isdigit():
                code_idx = idx
                break

        # A course row has: Course Code, Course Name..., Credits, Grade
        if code_idx != -1 and len(tokens) >= code_idx + 4:
            course_code = tokens[code_idx]
            grade = tokens[-1]
            credits_str = tokens[-2]
            course_name = " ".join(tokens[code_idx + 1 : -2])

            # Ensure credits is numeric and grade is a concise grade token (not text)
            if re.match(r"^\d+(?:\.\d+)?$", credits_str) and len(grade) <= 5:
                credits_val = int(float(credits_str)) if float(credits_str).is_integer() else float(credits_str)
                records.append((course_code, course_name, credits_val, grade))

    return records


def extract_courses_from_pdf(pdf_path: str) -> list[dict]:
    """
    Extract course details (Name, Roll No, Semester, Result Date, Course Code,
    Course Name, Credits, Grade) dynamically from a given PDF, with automatic
    pytesseract OCR fallback for scanned or image-based PDFs.
    """
    doc = pymupdf.open(pdf_path)
    full_text = "\n".join(page.get_text() for page in doc)

    # Check if direct text extraction returned text
    is_scanned = len(full_text.strip()) == 0

    student_name, roll_no = extract_student_info(full_text)
    semester = extract_semester_info(full_text)
    result_date = extract_result_date(full_text)

    records = []

    # First pass: try direct vector text extraction via PyMuPDF
    if not is_scanned:
        for page in doc:
            words = page.get_text("words")
            if not words:
                continue
            rows = parse_course_rows_from_words(words, line_tolerance=4.0)
            for course_code, course_name, credits, grade in rows:
                records.append({
                    "Name": student_name,
                    "Roll No": roll_no,
                    "Semester": semester,
                    "Result Date": result_date,
                    "Course Code": course_code,
                    "Course Name": course_name,
                    "No. of Credits": credits,
                    "Grade": grade
                })

    # Second pass: If scanned or no records found on document, use pytesseract OCR fallback
    if len(records) == 0 and len(doc) > 0:
        print(f"  [OCR Fallback] Invoking pytesseract for {os.path.basename(pdf_path)}...")
        ocr_texts = []
        ocr_records = []

        for page in doc:
            ocr_text, ocr_words = extract_page_words_with_ocr(page)
            if ocr_text:
                ocr_texts.append(ocr_text)
            if ocr_words:
                rows = parse_course_rows_from_words(ocr_words, line_tolerance=10.0)
                ocr_records.extend(rows)

        # If OCR extracted text, dynamically retrieve metadata
        if ocr_texts:
            combined_ocr_text = "\n".join(ocr_texts)
            if student_name == "Unknown" or roll_no == "Unknown":
                ocr_name, ocr_roll = extract_student_info(combined_ocr_text)
                if ocr_name != "Unknown":
                    student_name = ocr_name
                if ocr_roll != "Unknown":
                    roll_no = ocr_roll
            if semester in ("Unknown Semester", "N/A"):
                semester = extract_semester_info(combined_ocr_text)
            if result_date == "N/A":
                result_date = extract_result_date(combined_ocr_text)

            for course_code, course_name, credits, grade in ocr_records:
                records.append({
                    "Name": student_name,
                    "Roll No": roll_no,
                    "Semester": semester,
                    "Result Date": result_date,
                    "Course Code": course_code,
                    "Course Name": course_name,
                    "No. of Credits": credits,
                    "Grade": grade
                })

    return records


def format_excel(excel_path: str):
    """
    Apply professional styling, auto-fit column widths,
    and enable filter dropdowns on the table header.
    """
    wb = openpyxl.load_workbook(excel_path)
    ws = wb.active

    header_fill = PatternFill(start_color="1F4E79", end_color="1F4E79", fill_type="solid")
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    center_align = Alignment(horizontal="center", vertical="center")
    left_align = Alignment(horizontal="left", vertical="center")
    thin_border = Border(
        left=Side(style="thin", color="D9D9D9"),
        right=Side(style="thin", color="D9D9D9"),
        top=Side(style="thin", color="D9D9D9"),
        bottom=Side(style="thin", color="D9D9D9")
    )

    # Style header row
    for col in range(1, ws.max_column + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center_align

    # Style data rows: Name (col 1), Semester (col 3), and Course Name (col 6) left-aligned; others centered
    for row in range(2, ws.max_row + 1):
        for col in range(1, ws.max_column + 1):
            cell = ws.cell(row=row, column=col)
            cell.border = thin_border
            if col in (1, 3, 6):  # Name, Semester/Conduction Heading, Course Name
                cell.alignment = left_align
            else:
                cell.alignment = center_align

    # Set column widths with extra padding
    for col in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)

    # Enable filter dropdowns on header row
    ws.auto_filter.ref = f"A1:{get_column_letter(ws.max_column)}{ws.max_row}"

    # Freeze header row
    ws.freeze_panes = "A2"
    wb.save(excel_path)


def main():
    # 1. Select PDF files via OS-native file dialog
    pdf_files = select_pdf_files()

    if not pdf_files:
        print("No files selected.")
        return

    print(f"\nProcessing {len(pdf_files)} PDF file(s)...")

    all_records = []
    for pdf_file in pdf_files:
        records = extract_courses_from_pdf(pdf_file)
        if records:
            name = records[0]["Name"]
            roll = records[0]["Roll No"]
            sem = records[0]["Semester"]
            print(f"  - {os.path.basename(pdf_file)}: [{roll}] {name} ({sem}) -> {len(records)} courses")
        else:
            print(f"  - {os.path.basename(pdf_file)}: 0 courses extracted")
        all_records.extend(records)

    if not all_records:
        print("No course records extracted from the selected files.")
        return

    # 3. Build DataFrame
    df = pd.DataFrame(all_records)
    total_raw_records = len(df)

    # 4. Deduplicate multiple entries of the same course for a student by latest result date
    df["_parsed_date"] = pd.to_datetime(df["Result Date"], dayfirst=True, errors="coerce")
    # Sort ascending by student, course, and announcement date so the latest result is last
    df = df.sort_values(by=["Roll No", "Course Code", "_parsed_date"], ascending=[True, True, True])

    # Keep only the latest entry per student and course
    deduped_df = df.drop_duplicates(subset=["Roll No", "Course Code"], keep="last").copy()
    deduped_df = deduped_df.drop(columns=["_parsed_date"])

    duplicates_removed = total_raw_records - len(deduped_df)
    if duplicates_removed > 0:
        print(f"\nDeduplication: Consolidated {duplicates_removed} older course entry/entries to retain only the latest result per course.")

    # Sort final DataFrame by Student and Course Code
    deduped_df = deduped_df.sort_values(by=["Roll No", "Course Code"], ascending=[True, True])

    # 5. Open native save file dialog to choose location and name
    output_excel = get_save_file_path("student_grades.xlsx")
    if not output_excel:
        print("Save cancelled. Excel file was not saved.")
        return

    # 6. Save to Excel using pandas
    deduped_df.to_excel(output_excel, index=False)

    # 7. Format the Excel workbook (including header auto-filters)
    format_excel(output_excel)

    num_students = deduped_df["Roll No"].nunique()
    print(f"\nSuccessfully extracted {len(deduped_df)} unique course records across {num_students} student(s).")
    print(f"Saved, formatted, and filter-enabled Excel file at: {output_excel}\n")
    print("Summary of extracted courses:")
    print(deduped_df.to_string(index=False))


if __name__ == "__main__":
    main()